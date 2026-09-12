# -*- coding: utf-8 -*-
"""Verify a benchmark PPTX: page count + native editability stats (recursive).

Usage: python verify_pptx.py <deck.pptx>
Hard rule: pictures must be 0 (screenshot-skin = fail); text frames must be > 0.
"""
import sys
from pptx import Presentation

path = sys.argv[1]
prs = Presentation(path)
stats = {"shapes": 0, "texts": 0, "pics": 0, "groups": 0, "freeform": 0}
sample = []


def walk(shapes):
    for sh in shapes:
        stats["shapes"] += 1
        st = str(sh.shape_type)
        if sh.shape_type == 6:
            stats["groups"] += 1
            walk(sh.shapes)
            continue
        if sh.shape_type == 13:
            stats["pics"] += 1
        if "FREEFORM" in st:
            stats["freeform"] += 1
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                stats["texts"] += 1
                if len(sample) < 6:
                    sample.append(t[:24])


for s in prs.slides:
    walk(s.shapes)
verdict = "PASS" if (stats["pics"] == 0 and stats["texts"] > 0) else "FAIL (screenshot-skin or no text)"
print("file:", path)
print("slides:", len(prs.slides))
print(stats)
print("sample texts:", sample)
print("verdict:", verdict)
